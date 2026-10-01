from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
PPT = ROOT / "doc" / "SmartDent_Presentacion_Avance_1.pptx"

REPLACEMENTS = {
    "34": "19",
    "34 pruebas automatizadas aprobadas.": "19 pruebas unitarias aprobadas.",
    "Cobertura principal: autenticación · roles · citas · disponibilidad · historias clínicas · reportes · OpenAPI":
        "Cobertura unitaria: registro · autenticación · CRUD de servicios · reservas y reglas de citas",
    "AuthSecurityIntegrationTests.java": "LoginServiceUnitTest.java",
    "Prueba clave 02 · Reserva sin conflictos": "Prueba clave 02 · Sesiones sin doble cobro",
    "Comprueba persistencia, visibilidad por rol y prevención de cruces":
        "Comprueba una regla central del tratamiento odontológico",
    "PACIENTE": "SESIÓN",
    "La nueva cita aparece en su panel.": "La siguiente reserva continúa el tratamiento.",
    "ODONTÓLOGO": "SALDO",
    "Solo el profesional asignado puede verla.": "Las sesiones restantes se actualizan.",
    "ADMIN": "COBRO",
    "La reserva aparece en la agenda global.": "La segunda sesión tiene precio cero.",
    "CitaServiceIntegrationTests.java · fragmento real simplificado":
        "CitaServiceUnitTest.java · prueba aislada con Mockito",
    "Prueba clave 03 · Historia clínica": "Prueba clave 03 · CRUD de servicios",
    "La atención actualiza el expediente y marca la cita como atendida":
        "Valida la creación y normalización sin conectarse a MySQL",
    "CONDICIÓN": "AISLAMIENTO",
    "La cita debe estar confirmada antes de atenderse.": "El repositorio se reemplaza por un mock.",
    "TRAZABILIDAD": "NORMALIZACIÓN",
    "Diagnóstico y tratamiento quedan vinculados.": "El código se guarda en mayúsculas y sin espacios.",
    "Al guardar, la cita cambia a ATENDIDA.": "Se verifican precio y operación de guardado.",
    "HistoriaClinicaIntegrationTests.java · fragmento real simplificado":
        "ServicioServiceUnitTest.java · prueba aislada con Mockito",
}

LOGIN_CODE = '''@Test
void iniciaSesionYDevuelveTokenJwt() {
    Usuario usuario = usuarioActivo("paciente@correo.com");
    when(usuarioRepository.findByEmailIgnoreCase(
        "paciente@correo.com")).thenReturn(Optional.of(usuario));
    when(jwtService.generarToken(usuario)).thenReturn("jwt-prueba");

    var response = service.iniciarSesion(
        new LoginRequest(" PACIENTE@CORREO.COM ", "Clave1234"));

    verify(authenticationManager).authenticate(any(
        UsernamePasswordAuthenticationToken.class));
    assertThat(response.token()).isEqualTo("jwt-prueba");
    assertThat(response.tokenType()).isEqualTo("Bearer");
}'''

CITA_CODE = '''@Test
void segundaSesionDelMismoTratamientoNoVuelveACobrar() {
    Cita anterior = new Cita();
    anterior.setEstado(CitaEstado.ATENDIDA);
    anterior.setTratamientoCodigo("TRA-PRUEBA");
    anterior.setNumeroSesion(1);
    anterior.setTotalSesiones(4);
    when(citaRepository
        .findByPaciente_IdAndServicio_IdOrderByCreadoEnDesc(1L, 3L))
        .thenReturn(List.of(anterior));

    var cita = service.reservar("paciente@correo.com",
        request(fechaHabil(3), LocalTime.of(9, 0)));

    assertThat(cita.numeroSesion()).isEqualTo(2);
    assertThat(cita.sesionesRestantes()).isEqualTo(2);
    assertThat(cita.precioPactado()).isZero();
}'''

SERVICIO_CODE = '''@Test
void creaServicioNormalizandoElCodigo() {
    when(repository.save(any(Servicio.class)))
        .thenAnswer(invocation -> invocation.getArgument(0));

    var creado = service.crear(
        request(" srv-nuevo ", true));

    assertThat(creado.codigo()).isEqualTo("SRV-NUEVO");
    assertThat(creado.precio())
        .isEqualByComparingTo("120.00");
    verify(repository).save(any(Servicio.class));
}'''

presentation = Presentation(PPT)

for slide in presentation.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        if text in REPLACEMENTS:
            shape.text_frame.paragraphs[0].text = REPLACEMENTS[text]

for slide_number, code in ((11, LOGIN_CODE), (12, CITA_CODE), (13, SERVICIO_CODE)):
    slide = presentation.slides[slide_number - 1]
    blocks = [shape for shape in slide.shapes if shape.has_text_frame and "@Test" in shape.text]
    if len(blocks) != 1:
        raise RuntimeError(f"Bloque de código no encontrado en diapositiva {slide_number}.")
    blocks[0].text_frame.paragraphs[0].text = code

presentation.save(PPT)
print(f"Presentación actualizada: {PPT}")
print(f"Diapositivas: {len(presentation.slides)}")
