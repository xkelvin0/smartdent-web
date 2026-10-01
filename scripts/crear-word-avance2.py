from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "doc" / "Avance_2_SmartDent.md"
OUTPUT = ROOT / "SmartDent_Avance_2.docx"


def set_document_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for style_name, size, color in [
        ("Title", 20, RGBColor(0, 32, 64)),
        ("Heading 1", 16, RGBColor(0, 32, 64)),
        ("Heading 2", 14, RGBColor(0, 32, 64)),
        ("Heading 3", 12, RGBColor(0, 32, 64)),
    ]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        header[idx].text = value.strip()
        for paragraph in header[idx].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value.strip()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        values = [cell.strip() for cell in raw.split("|")]
        if not all(set(cell.replace(":", "").strip()) <= {"-"} for cell in values):
            rows.append(values)
        index += 1
    return rows, index


def add_code_block(document: Document, code_lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    for line in code_lines:
        run = paragraph.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def build_document() -> None:
    document = Document()
    set_document_style(document)

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SmartDent - Avance 2")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0, 32, 64)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Modelo de datos, JPA/Hibernate, CRUD, JPQL, transacciones, Spring Security y JWT")

    document.add_paragraph()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()

        if not line:
            index += 1
            continue

        if line.startswith("```"):
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            document.add_paragraph()
            continue

        if line.startswith("# "):
            # The custom cover already contains the main title.
            index += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            document.add_heading(line[5:].strip(), level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line.replace("`", "").strip())
        index += 1

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
